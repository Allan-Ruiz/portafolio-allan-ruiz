import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def style_header(ws, row_num, max_col, fill_hex="1E3A8A"):
    header_fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[row_num].height = 26

def auto_fit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if val.startswith("="):
                val = "12345678901234"
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 48)

def create_ambientes_limpios_excel(output_path):
    wb = openpyxl.Workbook()
    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )

    # ---------------------------------------------------------
    # HOJA 1: 1_Liquidacion_Transportistas
    # ---------------------------------------------------------
    ws1 = wb.active
    ws1.title = "1_Liquidacion_Transportistas"
    ws1["A1"] = "AMBIENTES LIMPIOS S.A. — CONTROL DE SALIDAS/VUELTAS DIARIAS Y LIQUIDACIÓN DE PAGO A TRANSPORTISTAS"
    ws1["A1"].font = Font(size=13, bold=True, color="065F46")
    ws1["A2"] = "Los camiones realizaban 1 o 2 vueltas al día (1 vuelta regreso 12:00-14:00 = $25.000 | 2 vueltas regreso 16:00-17:00 = $55.000). Se cuadraba con los respaldos impresos para firma del jefe y liberación del pago."
    ws1["A2"].font = Font(size=10, italic=True, color="4B5563")

    ws1["A4"] = "REGLA DE LIQUIDACIÓN DIARIA"
    ws1["B4"] = "MONTO A PAGAR"
    style_header(ws1, 4, 2, "065F46")
    ws1["A5"] = "1 Vuelta al día (Regreso aprox. 12:00 a 14:00 hrs)"
    ws1["B5"] = 25000
    ws1["B5"].number_format = "$#,##0"
    ws1["A6"] = "2 Vueltas al día (Regreso aprox. 16:00 a 17:00 hrs)"
    ws1["B6"] = 55000
    ws1["B6"].number_format = "$#,##0"

    headers1 = [
        "Fecha",
        "Transportista",
        "Patente_Camion",
        "Hora_Salida_1",
        "Hora_Vuelta_1 (12:00-14:00)",
        "Hora_Salida_2",
        "Hora_Vuelta_2 (16:00-17:00)",
        "Total_Vueltas_Dia (1 o 2)",
        "Monto_Pagar_Dia_CLP (Fórmula)",
        "Cuadratura_Archivo_Impreso_y_Firma_Jefe"
    ]
    ws1.append([])
    ws1.append(headers1)
    style_header(ws1, 8, len(headers1), "1E3A8A")

    registros = [
        ["2023-05-02", "Carlos Muñoz", "LJDK-42", "08:30", "12:45", "-", "-", 1, "Cuadrado con respaldo impreso -> Firmado por Jefe para pago"],
        ["2023-05-02", "Roberto Soto", "PPRT-18", "08:15", "12:30", "13:30", "16:40", 2, "Cuadrado con respaldo impreso -> Firmado por Jefe para pago"],
        ["2023-05-03", "Carlos Muñoz", "LJDK-42", "08:40", "13:10", "14:00", "16:55", 2, "Cuadrado con respaldo impreso -> Firmado por Jefe para pago"],
        ["2023-05-03", "Patricio Díaz", "KLMN-88", "09:00", "13:45", "-", "-", 1, "Cuadrado con respaldo impreso -> Firmado por Jefe para pago"],
        ["2023-05-04", "Roberto Soto", "PPRT-18", "08:20", "12:50", "13:45", "16:30", 2, "Cuadrado con respaldo impreso -> Firmado por Jefe para pago"],
        ["2023-05-04", "Carlos Muñoz", "LJDK-42", "08:35", "13:20", "-", "-", 1, "Cuadrado con respaldo impreso -> Firmado por Jefe para pago"],
        ["2023-05-05", "Patricio Díaz", "KLMN-88", "08:50", "12:55", "13:50", "16:45", 2, "Cuadrado con respaldo impreso -> Firmado por Jefe para pago"],
        ["2023-05-05", "Roberto Soto", "PPRT-18", "08:25", "13:15", "-", "-", 1, "Cuadrado con respaldo impreso -> Firmado por Jefe para pago"],
    ]

    for r_i, reg in enumerate(registros, start=9):
        f_pago = f"=IF(H{r_i}=1,$B$5,IF(H{r_i}=2,$B$6,0))"
        row_vals = reg[:8] + [f_pago, reg[8]]
        ws1.append(row_vals)
        for c_i in range(1, len(headers1) + 1):
            cell = ws1.cell(row=r_i, column=c_i)
            cell.border = thin_border
            if c_i == 9:
                cell.number_format = "$#,##0"

    tot_r = len(registros) + 9
    ws1.cell(row=tot_r, column=8, value="TOTAL LIQUIDACIÓN CUADRADA PARA PAGO").font = Font(bold=True)
    ws1.cell(row=tot_r, column=9, value=f"=SUM(I9:I{tot_r-1})").font = Font(bold=True)
    ws1.cell(row=tot_r, column=9).number_format = "$#,##0"

    # ---------------------------------------------------------
    # HOJA 2: 2_Conciliacion_Fisico_vs_ERP
    # ---------------------------------------------------------
    ws2 = wb.create_sheet("2_Conciliacion_Fisico_vs_ERP")
    ws2["A1"] = "CONTEOS CÍCLICOS Y CONCILIACIÓN BODEGA FÍSICA VS. SISTEMA ERP"
    ws2["A1"].font = Font(size=13, bold=True, color="1E3A8A")

    headers2 = ["Codigo_ERP", "Descripcion_Insumo_Bodega", "Unidad", "Saldo_Sistema_ERP", "Conteo_Fisico_Bodega", "Diferencia_Unidades (Fórmula)", "Costo_Unitario_CLP", "Valor_Diferencia_CLP (Fórmula)", "Estado_Conciliacion_ERP (Fórmula)"]
    ws2.append([])
    ws2.append(headers2)
    style_header(ws2, 3, len(headers2), "1E3A8A")

    erp_rows = [
        ["AL-10045", "Detergente Industrial Desengrasante Bidón 20L", "UN", 120, 120, 18500],
        ["AL-10082", "Desinfectante Amonio Cuaternario Concentrado 5L", "UN", 85, 82, 9400],
        ["AL-20110", "Rollo Papel Toalla Industrial Alta Absorción 2x300m", "PAQ", 240, 240, 6200],
        ["AL-30412", "Guante Nitrilo Industrial Caja 100 Unidades Talla L", "CAJ", 150, 146, 5800],
    ]

    for r_i, item in enumerate(erp_rows, start=4):
        f_dif = f"=E{r_i}-D{r_i}"
        f_val = f"=F{r_i}*G{r_i}"
        f_est = f'=IF(F{r_i}=0,"CUADRADO CON ERP","DISCREPANCIA - REVISAR MERMA O GUÍA")'
        ws2.append([item[0], item[1], item[2], item[3], item[4], f_dif, item[5], f_val, f_est])
        for c_i in range(1, len(headers2) + 1):
            cell = ws2.cell(row=r_i, column=c_i)
            cell.border = thin_border
            if c_i in (7, 8):
                cell.number_format = "$#,##0"

    for sheet in wb.worksheets:
        auto_fit_columns(sheet)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)

if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.abspath(__file__))
    p2 = os.path.join(base_dir, "02_ambientes_limpios_procesos_erp", "Caso02_AmbientesLimpios_Procesos_y_ERP.xlsx")
    create_ambientes_limpios_excel(p2)
    print("AMBIENTES_LIMPIOS_EXCEL_UPDATED_SUCCESSFULLY")
