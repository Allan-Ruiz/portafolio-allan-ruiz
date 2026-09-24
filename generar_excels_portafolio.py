import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def style_header(ws, row_num, max_col, fill_hex="1E3A8A"):
    header_fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[row_num].height = 26

def auto_fit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if val.startswith("="):
                val = "12345678901234"
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 48)

def create_bigpen_excel(output_path):
    wb = openpyxl.Workbook()
    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )
    alert_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    alert_font = Font(name="Calibri", size=10, bold=True, color="991B1B")
    ok_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    ok_font = Font(name="Calibri", size=10, bold=True, color="166534")

    # ---------------------------------------------------------
    # HOJA 1: 1_Validacion_SKU_y_WSP
    # ---------------------------------------------------------
    ws1 = wb.active
    ws1.title = "1_Validacion_SKU_y_WSP"
    ws1["A1"] = "CASO 01 LIBRERÍAS BIGPEN — PASO 1: VERIFICACIÓN DE SKU (GEMINI PRO + CLAUDE), CRUCE BUSCARX Y SUPERVISIÓN CON CLIENTE"
    ws1["A1"].font = Font(size=13, bold=True, color="1E3A8A")
    ws1["A2"] = "Cuando un SKU de la planilla no correspondía al producto real, se supervisaba y consultaba por WhatsApp al cliente para autorizar su corrección."
    ws1["A2"].font = Font(size=10, italic=True, color="4B5563")

    headers1 = [
        "Titulo_Original_Planilla",
        "SKU_Original_Cliente",
        "SKU_Cotejo_Gemini_Pro",
        "SKU_Cotejo_Claude",
        "Consistencia_IA (Fórmula)",
        "Supervision_BUSCARX (Fórmula)",
        "SKU_Validado_Final (Fórmula)",
        "Resolucion_WhatsApp_Cliente"
    ]
    ws1.append([])
    ws1.append(headers1)
    style_header(ws1, 4, len(headers1), "1E3A8A")

    qa_data = [
        ["  lapiz grafito hb n2 faber   ", "7806500112014", "7806500112014", "7806500112014", "SKU correcto -> Pasa a corrección de título, descripción web e imágenes"],
        ["marcador sharpie negro p. fina", "4005401125006", "071641300019", "071641300019", "Anomalía #1 consultada por WSP al cliente -> Autoriza corregir SKU"],
        ["cuaderno univ. torre 100h 7mm", "7807210045112", "7807210045112", "7807210045112", "SKU correcto -> Pasa a corrección de título, descripción web e imágenes"],
        ["destacador stabilo pastel amarillo", "8901180512099", "4006381492447", "4006381492447", "Anomalía #2 consultada por WSP al cliente -> Autoriza corregir SKU"],
        ["resma fotocopia carta 75g 500h", "SIN_SKU", "7806505001016", "7806505001016", "SKU obtenido con IA + validado con cliente -> Autoriza incorporar"],
        ["corrector bic shake squeeze 8ml", "3086123499110", "070330505926", "070330505926", "Anomalía #3 consultada por WSP al cliente -> Autoriza corregir SKU"],
        ["pegamento barra pritt 42g", "7804612008115", "7804612008115", "7804612008115", "SKU correcto -> Pasa a corrección de título, descripción web e imágenes"],
        ["boligrafo pilot g2 gel 0.7 azul", "4902505110044", "4902505163118", "4902505163118", "Anomalía #4 consultada por WSP al cliente -> Autoriza corregir SKU"],
        ["notas post-it 3m 76x76 amarillo", "6921734900118", "021200503351", "021200503351", "Anomalía #5 consultada por WSP al cliente -> Autoriza corregir SKU"],
    ]

    for idx, item in enumerate(qa_data, start=5):
        f_consistencia = f'=IF(C{idx}=D{idx},"100% CONSISTENTE","REVISAR")'
        f_dictamen = f'=IF(B{idx}=C{idx},"OK - SKU VALIDADO","ALERTA: CONSULTAR CLIENTE WSP")'
        f_sku_final = f'=IF(C{idx}=D{idx},C{idx},B{idx})'
        row = [item[0], item[1], item[2], item[3], f_consistencia, f_dictamen, f_sku_final, item[4]]
        ws1.append(row)
        is_anomaly = item[1] != item[2]
        for c_idx in range(1, len(headers1) + 1):
            cell = ws1.cell(row=idx, column=c_idx)
            cell.border = thin_border
            if is_anomaly and c_idx in (2, 6, 8):
                cell.fill = alert_fill
                cell.font = alert_font
            elif not is_anomaly and c_idx == 6:
                cell.fill = ok_fill
                cell.font = ok_font

    # ---------------------------------------------------------
    # HOJA 2: 2_Planilla_Importacion_Bsale
    # ---------------------------------------------------------
    ws2 = wb.create_sheet("2_Planilla_Importacion_Bsale")
    ws2["A1"] = "PASO 2: PLANILLA EXCEL DE IMPORTACIÓN A BSALE POS — TÍTULOS CORREGIDOS Y DESCRIPCIONES WEB COMPLETAS"
    ws2["A1"].font = Font(size=13, bold=True, color="065F46")
    ws2["A2"] = "Nota: Esta tabla se importa en Excel a Bsale. Las imágenes (mín. 2 por producto) se cargan por separado subiendo la carpeta con nombres SKU_1, SKU_2 (máx. SKU_8)."
    ws2["A2"].font = Font(size=10, italic=True, color="4B5563")

    headers2 = [
        "SKU_Validado_Bsale",
        "Titulo_Corregido_Bsale",
        "Descripcion_Web_Completa",
        "Categoria_Bsale",
        "Stock_Inicial",
        "Costo_Neto_CLP",
        "Precio_Venta_CLP",
        "Valorizacion_Costo_Total_CLP"
    ]
    ws2.append([])
    ws2.append(headers2)
    style_header(ws2, 4, len(headers2), "065F46")

    bsale_items = [
        [
            "7806500112014",
            "Lápiz Grafito HB N°2 Faber-Castell Hexagonal",
            "Lápiz grafito graduación HB N°2 con mina resistente al quiebre pegada con sistema SV. Ideal para escritura escolar, dibujo técnico y uso diario en oficina.",
            "ESCRITURA Y GRAFITO",
            140,
            190,
            390
        ],
        [
            "071641300019",
            "Marcador Permanente Sharpie Punta Fina Negro",
            "Marcador permanente de secado rápido resistente al agua y a la decoloración. Punta fina de alta precisión apta para papel, plástico, metal y vidrio.",
            "MARCADORES",
            85,
            650,
            1290
        ],
        [
            "7807210045112",
            "Cuaderno Universitario Torre 100 Hojas Cuadriculado 7mm",
            "Cuaderno universitario espiral doble recubierto, 100 hojas de papel 60g cuadriculado 7mm con tapa extra dura plastificada.",
            "CUADERNOS Y PAPELERIA",
            210,
            1150,
            2190
        ],
        [
            "4006381492447",
            "Destacador Stabilo Boss Original Pastel Amarillo",
            "Destacador tono pastel suave con tecnología Anti-Dry-Out (hasta 4 horas destapado sin secarse). Punta biselada para doble ancho de trazo (2 mm y 5 mm).",
            "DESTACADORES",
            64,
            890,
            1690
        ],
        [
            "070330505926",
            "Corrector Líquido BIC Shake 'n Squeeze 8 ml",
            "Corrector tipo lápiz con punta metálica de aguja para correcciones precisas y cuerpo blando de flujo controlado. Secado rápido y excelente cobertura.",
            "CORRECTORES",
            52,
            720,
            1390
        ],
        [
            "4902505163118",
            "Bolígrafo Pilot G-2 Gel Retráctil 0.7mm Azul",
            "Bolígrafo de tinta gel súper suave con grip ergonómico de caucho y clip retráctil. Trazo continuo de 0.7mm ideal para escritura intensiva.",
            "ESCRITURA Y GRAFITO",
            120,
            1190,
            2190
        ],
    ]

    for r_i, item in enumerate(bsale_items, start=5):
        val_formula = f"=E{r_i}*F{r_i}"
        ws2.append(item + [val_formula])
        for c_idx in range(1, len(headers2) + 1):
            cell = ws2.cell(row=r_i, column=c_idx)
            cell.border = thin_border
            if c_idx in (6, 7, 8):
                cell.number_format = "$#,##0"

    # ---------------------------------------------------------
    # HOJA 3: 3_Resumen_KPIs_Jefatura
    # ---------------------------------------------------------
    ws3 = wb.create_sheet("3_Resumen_KPIs_Jefatura")
    ws3["A1"] = "REPORTE EJECUTIVO DE MIGRACIÓN Y CATÁLOGO DIGITAL — LIBRERÍAS BIGPEN"
    ws3["A1"].font = Font(size=13, bold=True, color="1E3A8A")
    ws3.append([])
    ws3.append(["KPI / Entregable Operacional", "Resultado Logrado", "Detalle Técnico"])
    style_header(ws3, 3, 3, "1E3A8A")

    kpis = [
        ["Registros Brutos Auditados", "1.200+ productos", "Planillas informales depuradas con Power Query y BUSCARX"],
        ["Catálogo Final Importado en Bsale POS", "800+ SKU validados", "Títulos corregidos + Descripciones web completas en plantilla Excel Bsale"],
        ["Activos Digitales (Imágenes de Producto)", "1.600+ imágenes (Mín. 2 por SKU)", "Búsqueda asistida por IA y descarga bajo formato SKU_1, SKU_2 (hasta SKU_8)"],
        ["Control de Calidad con Cliente (WSP)", "5 anomalías críticas corregidas", "Supervisión humana cuando un SKU no correspondía al producto real"],
        ["Velocidad de Procesamiento Logístico", "100 productos / hora", "0% pérdida de información en la puesta en marcha de Bsale POS"],
    ]
    for r_i, k in enumerate(kpis, start=4):
        ws3.append(k)
        for c_i in range(1, 4):
            ws3.cell(row=r_i, column=c_i).border = thin_border

    for sheet in wb.worksheets:
        auto_fit_columns(sheet)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)

if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.abspath(__file__))
    p1 = os.path.join(base_dir, "01_bigpen_etl_bsale_pos", "Caso01_BigPEN_Pipeline_ETL_Bsale_POS.xlsx")
    create_bigpen_excel(p1)
    print("BIGPEN_EXCEL_UPDATED_SUCCESSFULLY")
